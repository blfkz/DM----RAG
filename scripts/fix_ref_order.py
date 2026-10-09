"""
scripts/fix_ref_order.py
修正参考文献后半段编号,使其与正文首次出现顺序一致:
正文顺序(段31 双端 → 段35 入库 → 段37 检索 → 段39 防幻觉 → 段46 42题):
  贾春燕+许娜芬=[18-19]、倪嘉辛=[20]、杨俊=[21]、黄红伟=[22]、李铂鑫=[23]、张力军+王昱婷=[24-25]
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

DOC = Path("docs/软件导刊投稿稿_湘西土特产智能助农问答系统设计与实现.docx")

# 正文标记替换(带方括号完整匹配;用 TMP 中转避免链式误改)
BODY_FIXES = [
    "[21-22]", "[18-19]",
    "[25]", "[TMPX]",
    "[20]", "[23]",
    "[19]", "[22]",
    "[18]", "[21]",
    "[TMPX]", "[20]",
    "[23-24]", "[24-25]",
]

# 文献条目编号重排(前缀替换)
REF_FIXES = [
    ("[18] YANG", "[TMPA] YANG"),
    ("[19] HUANG", "[TMPB] HUANG"),
    ("[20] LI B X", "[TMPC] LI B X"),
    ("[21] JIA", "[TMPD] JIA"),
    ("[22] XU N F", "[TMPE] XU N F"),
    ("[23] ZHANG L J", "[TMPF] ZHANG L J"),
    ("[24] WANG Y T", "[TMPG] WANG Y T"),
    ("[25] NI J X", "[TMPH] NI J X"),
    ("[18] 杨俊", "[TMPA] 杨俊"),
    ("[19] 黄红伟", "[TMPB] 黄红伟"),
    ("[20] 李铂鑫", "[TMPC] 李铂鑫"),
    ("[21] 贾春燕", "[TMPD] 贾春燕"),
    ("[22] 许娜芬", "[TMPE] 许娜芬"),
    ("[23] 张力军", "[TMPF] 张力军"),
    ("[24] 王昱婷", "[TMPG] 王昱婷"),
    ("[25] 倪嘉辛", "[TMPH] 倪嘉辛"),
    # 临时 -> 目标
    ("[TMPD]", "[18]"),  # 贾春燕
    ("[TMPE]", "[19]"),  # 许娜芬
    ("[TMPH]", "[20]"),  # 倪嘉辛
    ("[TMPA]", "[21]"),  # 杨俊
    ("[TMPB]", "[22]"),  # 黄红伟
    ("[TMPC]", "[23]"),  # 李铂鑫
    ("[TMPF]", "[24]"),  # 张力军
    ("[TMPG]", "[25]"),  # 王昱婷
]

MARK_RE = re.compile(r"\[\d+(?:-\d+)?\]")


def set_run_font(run, sup=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if sup:
        run.font.superscript = True


def rebuild_para(p, text):
    for run in list(p.runs):
        run._element.getparent().remove(run._element)
    pos = 0
    for m in MARK_RE.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            set_run_font(r)
        r = p.add_run(m.group(0))
        set_run_font(r, sup=True)
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        set_run_font(r)


def main():
    doc = Document(str(DOC))

    # 1) 正文标记
    n_body = 0
    for p in doc.paragraphs:
        t = p.text
        new_t = t
        for i in range(0, len(BODY_FIXES), 2):
            old, new = BODY_FIXES[i], BODY_FIXES[i + 1]
            if old in new_t:
                new_t = new_t.replace(old, new)
        if new_t != t:
            rebuild_para(p, new_t)
            n_body += 1
    print(f"[正文] 修正段落: {n_body}")

    # 2) 文献条目前缀
    n_ref = 0
    for p in doc.paragraphs:
        t = p.text
        if not t.strip().startswith("["):
            continue
        new_t = t
        for old, new in REF_FIXES:
            if new_t.startswith(old):
                new_t = new.replace(old, new, 1)
                break
        if new_t != t:
            for run in p.runs:
                run.text = ""
            p.runs[0].text = new_t
            for run in p.runs[1:]:
                run._element.getparent().remove(run._element)
            n_ref += 1
    print(f"[文献] 修正条目: {n_ref}")

    doc.save(str(DOC))
    print("[完成] 已保存")

    # 验证
    doc2 = Document(str(DOC))
    refs = {}
    for p in doc2.paragraphs:
        t = p.text.strip()
        m = re.match(r"^\[(\d+)\]\s", t)
        if m:
            refs[int(m.group(1))] = t[:40]
    print("[验证] 文献编号序列:", sorted(refs.keys()))
    # 正文顺序检查
    order = []
    for p in doc2.paragraphs:
        t = p.text
        if re.match(r"^\[\d+\]", t.strip()):
            break
        for m in re.finditer(r"\[(\d+)(?:-(\d+))?\]", t):
            order.append(int(m.group(1)))
    print("[验证] 正文标记顺序:", order)


if __name__ == "__main__":
    main()
