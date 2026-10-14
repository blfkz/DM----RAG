"""
scripts/make_kc_docx.py
生成《湖南省社科基金 2026 课题指南·第165/166项方向调研》Word 文档(发老师用)。
166(湖南古籍数字资源建设与智能整理)为主,165(AI失实信息)为交叉侧。
"""
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

OUT = Path("docs/课题方向调研_发老师版.docx")
SONG, HEI = "宋体", "黑体"
DARK = RGBColor(0x1F, 0x2D, 0x50)


def _cn(run, name=SONG, size=10.5, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = color


def para(doc, text, name=SONG, size=10.5, bold=False, color=None,
         align=None, indent=True, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _cn(r, name, size, bold, color)
    if align is not None:
        p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.space_after = Pt(space_after)
    return p


def h1(doc, text):
    return para(doc, text, name=HEI, size=14, bold=True, color=DARK,
                indent=False, space_after=8)


def h2(doc, text):
    return para(doc, text, name=HEI, size=12, bold=True, color=DARK,
                indent=False, space_after=6)


def body(doc, text):
    return para(doc, text)


def main():
    doc = Document()
    para(doc, "湖南省社会科学基金项目 2026 年度课题指南", name=HEI, size=16,
         bold=True, color=DARK, align=1, indent=False, space_after=2)
    para(doc, "第 166、165 项课题方向调研", name=HEI, size=15,
         bold=True, color=DARK, align=1, indent=False, space_after=10)

    # ============ 一、方向 166(主) ============
    h1(doc, "一、方向 166:湖南古籍数字资源建设与智能整理研究(重点方向)")
    h2(doc, "1. 方向内涵")
    body(doc, "本方向面向湖南古籍(简牍、帛书、方志、湖湘文献等)的数字化建设与智能化整理,"
              "涵盖古籍数字化采集、文本结构化、自动句读、知识图谱构建、AI 辅助释读与智能检索服务等环节,"
              "是湖南省实施“文化+科技”融合战略、建设文化强省的基础性工作。")
    h2(doc, "2. 政策背景")
    body(doc, "2024 年 11 月,习近平总书记考察云梦县博物馆时指出:“古代简牍非常珍贵,"
              "是我们国家信史的重要实物佐证,要善加保护,做好研究”,为简牍数字化保护研究提供了根本遵循。"
              "湖南省委、省政府提出“文化+科技”融合发展部署,省委书记沈晓明同志两次就里耶秦简保护研究利用作出批示;"
              "2022 年起实施“湖湘文物数字化工程”与“湖南简帛整理保护工程”,"
              "并将“推进里耶秦简数字化采集与数据库建设”列入重点工作。")
    h2(doc, "3. 湖南省已有建设进展")
    body(doc, "(1) 湖南省古籍智慧化服务平台:由湖南图书馆于 2025 年建成上线,集古籍书目数据库、"
              "在线阅览、AI 句读、AI 翻译、AI 质检于一体,已上线古籍影像 10 万页、书目数据 3.6 万余条,"
              "累计访问量超 32 万人次。(2)《湖湘文库》AI 版:湖南出版集团 2025 年发布,基于 702 卷、"
              "约 4 亿字湖湘文献,融合大语言模型与知识图谱,提供 AI 辅助阅读与知识问答服务。")
    h2(doc, "4. 里耶秦简数智化现状与需求(核心场景)")
    body(doc, "里耶秦简出土于湘西龙山县里耶古城,共约 3.8 万枚,有字简约 1.7 万枚,"
              "是研究秦代政治制度、基层行政的第一手珍贵史料。目前:(1) 简牍数字化采集尚未完成,"
              "湘西州政协已专题协商“加快推进数字化采集和数据库建设”;(2) 有字简释读任务繁重,"
              "传统人工释读效率低、周期长;(3) 简牍图像存在残损、字迹模糊、异体字多等难题,"
              "通用古籍 OCR 准确率仅 70%~80%;(4) 通用大模型古籍知识不足,AI 辅助释读与整理"
              "存在输出失实信息(幻觉)的风险,亟需建立面向简牍场景的失实信息抑制机制。")
    h2(doc, "5. 学校前期工作基础")
    body(doc, "吉首大学于 2025 年 2 月成立里耶秦简研究院(现有教授 23 人、博士研究人员 42 人);"
              "2026 年 4 月,省教育厅、省文旅厅与吉首大学共同研究推进里耶秦简研究院建设,"
              "目标建成国家级秦简研究中心,并与湖南简帛研究中心共建联合在线研究平台。"
              "吉首大学学报已刊发《简牍数智化技术进展与展望:里耶秦简的数智化探索》,"
              "系统梳理了简牍图像增强、字符分割、字符识别与复原等技术的进展与挑战,"
              "指出当前存在“标准缺失、图像质量低”等瓶颈,为课题申报奠定了前期基础。")

    # ============ 二、方向 165(交叉侧) ============
    h1(doc, "二、方向 165:人工智能生成失实信息传播风险研究(交叉支撑方向)")
    h2(doc, "1. 方向内涵")
    body(doc, "本方向研究生成式人工智能(大模型)所产生失实信息(即“幻觉”)的生成机理、"
              "传播规律与风险治理,核心包括失实信息的检测、抑制方法与监管框架。")
    h2(doc, "2. 研究现状")
    body(doc, "国际研究在健康、政治等大众传播领域已较为活跃:2025—2026 年有多篇系统综述指出,"
              "生成式 AI 显著提高了虚假信息的产量、传播速度与感知可信度,"
              "用户区分 AI 与人类虚假信息的能力接近随机水平,检测与生成进入对抗性军备竞赛;"
              "治理策略主张按领域分层。中文研究方面,刘泽垣等《大语言模型的幻觉问题研究综述》"
              "(软件学报,2025)与袁乐等《大语言模型检索增强生成优化技术研究综述》"
              "(计算机学报,2026)系统梳理了幻觉的成因、评估与缓解技术。")
    body(doc, "总体判断:现有研究集中于大众传播场景,面向古籍整理等专业领域知识服务场景的"
              "人工智能失实信息抑制研究几乎空白,本课题恰好填补该交叉空白。")

    # ============ 三、交叉点 ============
    h1(doc, "三、两方向交叉点与课题切入点")
    body(doc, "以方向 166 为主场景、方向 165 为核心问题,课题切入点为:"
              "面向湖南古籍(里耶秦简)智能整理场景的人工智能失实信息抑制方法研究。"
              "具体思路:依托里耶秦简释文数据构建知识库,研究古籍领域问答场景下"
              "大模型失实信息的抑制方法(检索增强生成、检索置信度评估、引用校验与拒答机制),"
              "并建立量化评估体系(以忠实度指标衡量,使失实率由基线水平显著下降至可接受阈值),"
              "为湖南古籍智能化服务平台与里耶秦简数字化建设提供技术支撑。")

    # ============ 四、文献 ============
    h1(doc, "四、主要参考文献")
    refs = [
        "[1] 刘泽垣,王鹏江,宋晓斌,等.大语言模型的幻觉问题研究综述[J].软件学报,2025,36(3):1152-1185.",
        "[2] 袁乐,刘绍华,王禹,等.大语言模型检索增强生成优化技术研究综述[J].计算机学报,2026,49(2):383-422.",
        "[3] Asai A, Wu Z Q, Wang Y Z, et al. Self-RAG: Learning to Retrieve, Generate, and Critique through Self-Reflection[C]//ICLR. 2024.",
        "[4] Yan S Q, Gui J C, Zhu Y T, et al. Corrective Retrieval Augmented Generation[R]. arXiv:2401.15884, 2024.",
        "[5] Es S, James J, Espinosa-Anke L, et al. RAGAS: Automated Evaluation of Retrieval Augmented Generation[R]. arXiv:2309.15217, 2023.",
        "[6] 简牍数智化技术进展与展望:里耶秦简的数智化探索[J].吉首大学学报(武陵山片区研究专栏).",
        "[7] Generative AI and health misinformation: production, propagation, and mitigation—a systematic review[J]. BMC Public Health, 2026, 26.",
        "[8] Generative AI and misinformation: a scoping review[J]. AI & Society, 2025.",
        "[9] 全球视野下的中国古籍数智化演进与未来趋势报告[R]. 世界互联网大会(WIC), 2026.",
    ]
    for r in refs:
        para(doc, r, size=9.5, indent=False, space_after=3)

    doc.save(str(OUT))
    print(f"[完成] {OUT}")


if __name__ == "__main__":
    main()
