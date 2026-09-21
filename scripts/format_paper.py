"""
scripts/format_paper.py
把用户修改后的论文初稿 Word 文档整理成普刊投稿格式:
- 标题黑体居中、作者楷体居中、单位宋体居中
- 摘要/关键词前缀加粗,正文宋体五号、首行缩进 2 字符、1.5 倍行距
- 一级标题黑体小四、二级标题黑体五号
- 修复:3.1 标题与正文粘连、补"5 实验测试"一级标题、正文误设为标题样式

输入:docs/论文初稿.docx(用户修改后的版本)
输出:docs/论文投稿稿.docx
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

SRC = Path("docs/论文初稿.docx")
OUT = Path("docs/论文投稿稿.docx")

HEI = "黑体"
KAI = "楷体"
SONG = "宋体"


def set_font(run, name, size, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def style_paragraph(p, font=SONG, size=10.5, bold=False, align=None, indent=True, spacing=1.5):
    p.style = p.part.document.styles["Normal"]
    for run in p.runs:
        set_font(run, font, size, bold)
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = spacing
    if indent:
        pf.first_line_indent = Pt(size * 2)  # 首行缩进 2 字符


def main():
    doc = Document(str(SRC))
    paras = doc.paragraphs

    for i, p in enumerate(paras):
        text = p.text.strip()
        if not text:
            continue

        # 1) 标题(第一段)
        if i == 0:
            style_paragraph(p, font=HEI, size=16, bold=True,
                            align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
            continue
        # 2) 作者行
        if i == 1:
            style_paragraph(p, font=KAI, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
            continue
        # 3) 单位行(含括号)
        if i == 2:
            style_paragraph(p, font=SONG, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
            continue

        # 4) 摘要/关键词行:前缀加粗
        if re.match(r"^(摘要|关键词|Abstract|Key words)[::]", text):
            style_paragraph(p, font=SONG, size=10.5)
            if p.runs and re.match(r"^(摘要|关键词|Abstract|Key words)", p.runs[0].text):
                set_font(p.runs[0], SONG, 10.5, bold=True)
            continue

        # 5) 参考文献标题
        if text == "参考文献":
            style_paragraph(p, font=HEI, size=12, bold=True, indent=False)
            continue

        # 6) 参考文献条目
        if re.match(r"^\[\d+\]", text):
            style_paragraph(p, font=SONG, size=10.5, indent=False)
            continue

        # 7) 一级标题(如 "1 引言"、"6 总结与展望")
        if re.match(r"^[1-9]\s+\S", text) and not re.match(r"^\d+\.\d+", text):
            style_paragraph(p, font=HEI, size=12, bold=True, indent=False)
            continue

        # 8) 二级标题(如 "2.1 大语言模型与云端 API 调用")
        if re.match(r"^\d+\.\d+", text):
            # 修复:标题与正文粘连(用户编辑时把 3.1 标题和正文合在了一段)
            if text.startswith("3.1 系统架构") and len(text) > 30:
                title = "3.1 系统架构"
                body = text[len(title):].strip()
                for run in p.runs:
                    run.text = ""
                if p.runs:
                    p.runs[0].text = title
                style_paragraph(p, font=HEI, size=10.5, bold=True, indent=False)
                new_p = p.insert_paragraph_before(body)
                new_p._element.addprevious(p._element)  # 调整顺序:标题在前、正文在后
                style_paragraph(new_p, font=SONG, size=10.5)
                print("[修复] 3.1 粘连段落已拆分")
                continue
            style_paragraph(p, font=HEI, size=10.5, bold=True, indent=False)
            continue

        # 9) 表标题
        if re.match(r"^表\s*\d", text):
            style_paragraph(p, font=SONG, size=10.5, bold=True,
                            align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
            continue

        # 10) 其余:正文
        style_paragraph(p, font=SONG, size=10.5)

    # 补"5 实验测试"一级标题(在 5.1 之前)
    for p in doc.paragraphs:
        if p.text.strip().startswith("5.1 实验环境"):
            new_p = p.insert_paragraph_before("5 实验测试")
            style_paragraph(new_p, font=HEI, size=12, bold=True, indent=False)
            print("[修复] 已补上'5 实验测试'一级标题")
            break

    doc.save(str(OUT))
    print(f"[完成] 投稿格式稿已生成:{OUT}")


if __name__ == "__main__":
    main()
