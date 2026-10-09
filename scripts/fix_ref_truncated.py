"""
scripts/fix_ref_truncated.py
重写被截断的 4 条文献(编号 18/19/24/25,英文行+中文行共 8 行)。
"""
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

DOC = Path("docs/软件导刊投稿稿_湘西土特产智能助农问答系统设计与实现.docx")

FULL = {
    "[18] JIA": ("[18] JIA C Y, FANG W J, XIE Y W, et al. Research on campus question answering "
                 "system supported by retrieval augmented generation technology[J]. "
                 "Journal on Communications, 2024, 45(Z2): 248-254."),
    "[18] 贾春燕": ("[18] 贾春燕,方伟杰,谢宇威,等.检索增强生成技术支持下的校园问答系统研究[J]."
                "通信学报,2024,45(Z2):248-254."),
    "[19] XU": ("[19] XU N F, HUANG Z L, XIE Z H. Intelligent e-commerce customer service system "
                "driven by DeepSeek[J]. Journal of Fujian Computer, 2025(12): 72-77."),
    "[19] 许娜芬": ("[19] 许娜芬,黄卓霖,谢中华.DeepSeek驱动的智能电商客户服务系统[J].福建电脑,"
                "2025(12):72-77."),
    "[24] ZHANG L J": ("[24] ZHANG L J, LIU S, LIAO J T, et al. Design and implementation of "
                       "question answering system for computer network experiment course based on "
                       "large model retrieval augmented generation[J]. Experimental Technology and "
                       "Management, 2024, 41(12): 186-192."),
    "[24] 张力军": ("[24] 张力军,刘偲,廖纪童,等.基于大模型检索增强生成的计算机网络实验课程问答系统"
                "设计与实现[J].实验技术与管理,2024,41(12):186-192."),
    "[25] WANG Y T": ("[25] WANG Y T, CHEN B, YAN Q, et al. Retrieval augmented generation question "
                      "answering based on problem-oriented prompt learning and multi-path "
                      "reasoning[J]. Computer Engineering and Applications, 2025, 61(12): 120-128."),
    "[25] 王昱婷": ("[25] 王昱婷,陈波,闫强,等.基于问题导向提示学习和多路推理的检索增强生成问答[J]."
                "计算机工程与应用,2025,61(12):120-128."),
}


def set_run_font(run):
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def main():
    doc = Document(str(DOC))
    fixed = 0
    for p in doc.paragraphs:
        t = p.text.strip()
        for key, full in FULL.items():
            if t.startswith(key) and t != full:
                for run in list(p.runs):
                    run._element.getparent().remove(run._element)
                r = p.add_run(full)
                set_run_font(r)
                fixed += 1
                print(f"[修] {full[:50]}...")
                break
    print(f"[完成] 修复 {fixed} 行")
    doc.save(str(DOC))

    # 终验:所有文献行长度
    doc2 = Document(str(DOC))
    started = False
    for p in doc2.paragraphs:
        t = p.text.strip()
        if t == "参考文献":
            started = True
            continue
        if not started or not t:
            continue
        if len(t) < 60:
            print(f"[异常短] {t}")


if __name__ == "__main__":
    main()
