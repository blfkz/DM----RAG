"""
scripts/make_plan_docs.py
把 docs/商业计划书.md 自动转成 Word(.docx)与 PDF 文件,用于大创网上传(要求 PDF <20M)。
改完计划书内容后重跑本脚本即可重新生成。

用法:.venv/Scripts/python scripts/make_plan_docs.py
"""
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

MD_PATH = Path("docs/商业计划书.md")
OUT_DOCX = Path("docs/商业计划书.docx")
OUT_PDF = Path("docs/商业计划书.pdf")


# ============ 简易 Markdown 解析(标题/段落/列表/表格/引用) ============

def _split_bold(text):
    """把 **粗体** 拆成 [(文字, 是否加粗), ...]"""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    out = []
    for p in parts:
        if p.startswith("**") and p.endswith("**"):
            out.append((p[2:-2], True))
        elif p:
            out.append((p, False))
    return out


def parse_md(text):
    """返回块列表:[("h1"|"h2"|"h3"|"p"|"quote"|"li"|"table", 内容)]"""
    blocks = []
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.strip() == "---":
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            blocks.append((f"h{len(m.group(1))}", m.group(2).strip()))
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            table = []
            for r in rows:
                cells = [c.strip() for c in r.strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    continue  # 表格分隔行
                table.append(cells)
            blocks.append(("table", table))
            continue
        if line.startswith("- "):
            items = []
            while i < len(lines) and lines[i].strip().startswith("- "):
                items.append(lines[i].strip()[2:])
                i += 1
            blocks.append(("li", items))
            continue
        if line.startswith(">"):
            blocks.append(("quote", line.lstrip("> ").strip()))
            i += 1
            continue
        blocks.append(("p", line))
        i += 1
    return blocks


# ============ 生成 Word ============

def make_docx(blocks):
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    for kind, content in blocks:
        if kind == "h1":
            doc.add_heading(content, level=1)
        elif kind == "h2":
            doc.add_heading(content, level=2)
        elif kind == "h3":
            doc.add_heading(content, level=3)
        elif kind == "p":
            p = doc.add_paragraph()
            for text, bold in _split_bold(content):
                run = p.add_run(text)
                run.bold = bold
        elif kind == "quote":
            p = doc.add_paragraph()
            run = p.add_run(content)
            run.italic = True
        elif kind == "li":
            for item in content:
                doc.add_paragraph(item, style="List Bullet")
        elif kind == "table":
            t = doc.add_table(rows=len(content), cols=len(content[0]))
            t.style = "Light Grid Accent 1"
            for ri, row in enumerate(content):
                for ci, cell in enumerate(row):
                    t.rows[ri].cells[ci].text = cell
    doc.save(OUT_DOCX)
    return OUT_DOCX


# ============ 生成 PDF ============

def _register_cn_font():
    """注册中文字体:优先黑体(单文件),其次微软雅黑/宋体(TTC 集合,用 subfontIndex 加载)。"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        ("simhei", r"C:\Windows\Fonts\simhei.ttf", "ttf"),
        ("msyh", r"C:\Windows\Fonts\msyh.ttc", "ttc"),
        ("simsun", r"C:\Windows\Fonts\simsun.ttc", "ttc"),
    ]
    for name, path, kind in candidates:
        try:
            if kind == "ttf":
                pdfmetrics.registerFont(TTFont("CN", path))
            else:
                pdfmetrics.registerFont(TTFont("CN", path, subfontIndex=0))
            print(f"[字体] 已加载:{path}")
            return
        except Exception as e:  # noqa: BLE001
            print(f"[字体] {path} 加载失败:{e}")
            continue
    raise RuntimeError("未找到可用中文字体,请检查 C:\\Windows\\Fonts")


def make_pdf(blocks):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    _register_cn_font()

    st_h1 = ParagraphStyle("h1", fontName="CN", fontSize=18, leading=26, spaceAfter=12,
                           textColor=colors.HexColor("#1a1a1a"))
    st_h2 = ParagraphStyle("h2", fontName="CN", fontSize=14, leading=22, spaceBefore=10, spaceAfter=6,
                           textColor=colors.HexColor("#2a2a2a"))
    st_h3 = ParagraphStyle("h3", fontName="CN", fontSize=12, leading=18, spaceBefore=8, spaceAfter=4)
    st_p = ParagraphStyle("p", fontName="CN", fontSize=10.5, leading=17, spaceAfter=6)
    st_quote = ParagraphStyle("quote", fontName="CN", fontSize=10, leading=16, spaceAfter=6,
                              textColor=colors.HexColor("#555555"), leftIndent=12)
    st_li = ParagraphStyle("li", fontName="CN", fontSize=10.5, leading=17, spaceAfter=3, leftIndent=14)
    st_cell = ParagraphStyle("cell", fontName="CN", fontSize=9, leading=13)
    st_cell_head = ParagraphStyle("cellhead", fontName="CN", fontSize=9, leading=13)

    def esc(text):
        return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story = []
    for kind, content in blocks:
        if kind == "h1":
            story.append(Paragraph(esc(content), st_h1))
        elif kind == "h2":
            story.append(Paragraph(esc(content), st_h2))
        elif kind == "h3":
            story.append(Paragraph(esc(content), st_h3))
        elif kind == "p":
            story.append(Paragraph("".join(f"<b>{esc(t)}</b>" if b else esc(t)
                                           for t, b in _split_bold(content)), st_p))
        elif kind == "quote":
            story.append(Paragraph(esc(content), st_quote))
        elif kind == "li":
            for item in content:
                story.append(Paragraph(f"• {esc(item)}", st_li))
        elif kind == "table":
            data = []
            for ri, row in enumerate(content):
                style = st_cell_head if ri == 0 else st_cell
                data.append([Paragraph(esc(c), style) for c in row])
            # 列宽均分页面可用宽度(A4 21cm - 左右边距各 2.2cm = 16.6cm),避免宽表格超出页面被截断
            n_cols = len(content[0])
            col_w = 16.6 / n_cols
            t = Table(data, colWidths=[col_w * cm] * n_cols)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#bbbbbb")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2f7")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(t)
            story.append(Spacer(1, 8))

    doc = SimpleDocTemplate(str(OUT_PDF), pagesize=A4,
                            leftMargin=2.2 * cm, rightMargin=2.2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm,
                            title="乡音智答——湘西土特产小店 AI 客服 商业计划书")
    doc.build(story)
    return OUT_PDF


def main():
    print("=" * 50)
    print("商业计划书 → Word + PDF 自动生成")
    print("=" * 50)
    text = MD_PATH.read_text(encoding="utf-8")
    blocks = parse_md(text)
    print(f"[解析] 共 {len(blocks)} 个内容块")
    docx = make_docx(blocks)
    pdf = make_pdf(blocks)
    print(f"[生成] Word:{docx}({docx.stat().st_size / 1024:.0f} KB)")
    print(f"[生成] PDF:{pdf}({pdf.stat().st_size / 1024:.0f} KB)")
    print(f"[提示] PDF 大小 {pdf.stat().st_size / 1024 / 1024:.1f} MB(大创网要求 <20M)")
    print("=" * 50)


if __name__ == "__main__":
    main()
