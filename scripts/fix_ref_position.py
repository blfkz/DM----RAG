"""
scripts/fix_ref_position.py
修复参考文献列表位置:文献行被插到"参考文献"标题之前且顺序倒置。
按编号升序排序,统一移到标题之后。
"""
import re
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

DOC = Path("docs/软件导刊投稿稿_湘西土特产智能助农问答系统设计与实现.docx")


def main():
    doc = Document(str(DOC))

    # 找参考文献标题
    title_p = None
    for p in doc.paragraphs:
        if p.text.strip() == "参考文献":
            title_p = p
            break
    assert title_p is not None, "未找到参考文献标题"

    # 收集标题之前、以 [数字] 开头的文献行
    ref_paras = []
    for p in doc.paragraphs:
        if p._element is title_p._element:
            break
        t = p.text.strip()
        m = re.match(r"^\[(\d+)\]", t)
        if m:
            ref_paras.append((int(m.group(1)), p))

    print(f"[收集] 标题前的文献行: {len(ref_paras)}")

    # 按编号排序(同一编号英文行在前)
    def sort_key(item):
        num, p = item
        en = p.text.strip().startswith(f"[{num}] ") and bool(re.match(rf"^\[{num}\]\s*[A-Z]", p.text.strip()))
        return (num, 0 if en else 1)

    ref_paras.sort(key=sort_key)

    # 移到标题之后,按序排列
    anchor = title_p._element
    for _, p in ref_paras:
        anchor.addnext(p._element)
        anchor = p._element

    doc.save(str(DOC))
    print("[完成] 已保存")

    # 验证:顺序打印
    doc2 = Document(str(DOC))
    started = False
    nums = []
    for p in doc2.paragraphs:
        t = p.text.strip()
        if t == "参考文献":
            started = True
            continue
        if not started:
            continue
        m = re.match(r"^\[(\d+)\]", t)
        if m:
            nums.append(int(m.group(1)))
    print("[验证] 标题后文献编号顺序:", nums)
    print("[验证] 共", len(nums), "行")


if __name__ == "__main__":
    main()
